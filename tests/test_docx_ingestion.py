import io
import pytest
import asyncio
from fastapi import HTTPException
from services.ingestion import DocumentIngestionService
import docx

def create_sample_docx_bytes(title: str = "Test Thesis", sections: dict = None) -> bytes:
    doc = docx.Document()
    doc.add_heading(title, level=0)
    
    if sections:
        for heading, body in sections.items():
            doc.add_heading(heading, level=1)
            doc.add_paragraph(body)
    else:
        doc.add_heading("Chapter 1: Introduction", level=1)
        doc.add_paragraph("This is the introduction section discussing background.")
        doc.add_heading("Chapter 2: Methodology", level=1)
        doc.add_paragraph("We analyzed N=100 participants using standard metrics.")
        
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@pytest.mark.asyncio
async def test_docx_text_extraction():
    docx_bytes = create_sample_docx_bytes()
    text = await DocumentIngestionService.extract_plaintext_from_upload(docx_bytes, "manuscript.docx")
    
    assert "Chapter 1: Introduction" in text
    assert "Chapter 2: Methodology" in text
    assert "N=100 participants" in text

@pytest.mark.asyncio
async def test_unsupported_format_raises_400():
    with pytest.raises(HTTPException) as exc_info:
        await DocumentIngestionService.extract_plaintext_from_upload(b"some plain text", "test.txt")
    assert exc_info.value.status_code == 400
    assert "Unsupported file format" in exc_info.value.detail

@pytest.mark.asyncio
async def test_empty_docx_raises_400():
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    
    with pytest.raises(HTTPException) as exc_info:
        await DocumentIngestionService.extract_plaintext_from_upload(buf.getvalue(), "empty.docx")
    assert exc_info.value.status_code == 400
    assert "appears to be empty" in exc_info.value.detail
